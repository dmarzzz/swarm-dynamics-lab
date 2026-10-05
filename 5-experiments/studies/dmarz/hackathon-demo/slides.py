"""Static slides of the four-minute cut, as a .pptx that Google Slides imports.

    python3 slides.py [_out/swarm-of-theseus-slides.pptx]

One slide per scene state: the scene is rendered at a chosen point (a still, no animation), the cube mark is laid on
bottom right as in the film, and the picture fills a 16:9 slide. The scene's narration goes in the speaker notes.
Needs python-pptx, ffmpeg and the same `kit/` folder the film needs.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

HERE = Path(__file__).parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / '_out' / 'swarm-of-theseus-slides.pptx'
FRAMES = HERE / '_out' / 'slides'

# (scene id, point in the scene 0..1, slide title for the outline and alt text, which part of the narration: None = all,
#  or (first sentence, last sentence) as a 0-based slice of that scene's sentences)
SLIDES = [
    ('team', 0.97, 'Swarm of Theseus', None),
    ('why', 0.97, 'The problem: the Hugging Face incident', None),
    ('frame', 0.33, 'What is distributional safety?', (0, 3)),
    ('frame', 0.99, 'Three agents, each safe on its own', (3, None)),
    ('questions', 0.97, 'Three questions for distributional safety', None),
    ('tracks', 0.97, '219 candidate hypotheses, 15 research areas', None),
    ('lab', 0.97, 'An open source, mostly-autonomous swarm research lab', None),
    ('stack', 0.99, 'Open source, infrastructure as code', None),
    ('highlights', 0.97, '118 studies, three highlighted', None),
    ('sybil', 0.74, 'Result 1: Thou shalt not split (the evasion)', (0, 4)),
    ('sybil', 0.995, 'Result 1: Thou shalt not split (one sentence added)', (4, None)),
    ('roadmap1', 0.97, 'What result 1 changes', None),
    ('result3', 0.97, 'Result 2: How to win agents and influence swarms', None),
    ('roadmap2', 0.97, 'What result 2 changes', None),
    ('theseus50', 0.99, 'Result 3: Swarm of Theseus', None),
    ('roadmap3', 0.97, 'What result 3 changes', None),
    ('more', 0.97, 'Ten more experiments', None),
    ('next', 0.52, 'Next steps', (0, 1)),
    ('next', 0.95, 'Swarm of Theseus: site and repository', (1, None)),
]

# the narration file is written for a speech synthesiser; put it back into reading form for the notes
READ = [('G P T six Sol', 'GPT-6 Sol'), ('A G I', 'AGI'), ('A I ', 'AI '), ('C I', 'CI'), ('M I T', 'MIT'), ('Open Tofu', 'OpenTofu'),
        ('swarm safety dot org', 'swarmsafety.org'), ('Single agent', 'Single-agent'), ('group level', 'group-level')]


def notes_for(narration, scene, part):
    text = narration.get(scene, '')
    for a, b in READ: text = text.replace(a, b)
    if part is None: return text
    sentences = re.split(r'(?<=[.?!])\s+', text)
    return ' '.join(sentences[part[0]:part[1]])


def frames():
    """Render each still with the film's own page, then lay the cube mark on it."""
    FRAMES.mkdir(parents=True, exist_ok=True)
    cube = FRAMES / '_cube.png'
    mark = json.loads(re.search(r"mark: (\{[^}]*\})", (HERE / 'index.html').read_text()).group(1).replace('x:', '"x":').replace('y:', '"y":').replace('size:', '"size":'))
    subprocess.run(['ffmpeg', '-y', '-v', 'error', '-c:v', 'libvpx-vp9', '-ss', '1', '-i', str(HERE / 'kit' / 'cube-loop-alpha.webm'), '-frames:v', '1',
                    '-vf', f"crop=380:380:0:0,scale={mark['size']}:{mark['size']}", str(cube)], check=True)
    paths = []
    for n, (scene, p, _, _) in enumerate(SLIDES, 1):
        subprocess.run(['node', 'shot.mjs', scene, str(p), 'film=long'], cwd=HERE, check=True, capture_output=True)
        out = FRAMES / f'{n:02d}-{scene}.png'
        subprocess.run(['ffmpeg', '-y', '-v', 'error', '-i', str(HERE / '_shots' / f'{scene}-{p}.png'), '-i', str(cube),
                        '-filter_complex', f"overlay={mark['x']}:{mark['y']}", str(out)], check=True)
        paths.append(out)
        print(out.name)
    return paths


def build(paths):
    narration = {s['scene']: s['text'] for s in json.loads((HERE / 'narration-long.json').read_text())['segments']}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for path, (scene, _, title, part) in zip(paths, SLIDES):
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(str(path), Emu(0), Emu(0), width=prs.slide_width, height=prs.slide_height)
        pic.name = title
        pic._element.nvPicPr.cNvPr.set('descr', title)             # alt text: what the slide shows
        slide.notes_slide.notes_text_frame.text = notes_for(narration, scene, part)
    prs.core_properties.title = 'Swarm of Theseus: Swarm Dynamics Lab'
    prs.core_properties.author = 'Swarm of Theseus'
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f'wrote {OUT}: {len(paths)} slides, {OUT.stat().st_size / 1e6:.1f} MB')


if __name__ == '__main__':
    build(frames())
